"""Извлечение таблиц из PDF в Excel.

Можно запускать и без веб-сервера:
    python -m app.converter file.pdf [file2.pdf ...] [-o excel/] [--single-sheet]
"""

import argparse
import io
import re
import sys
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

import pymupdf
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

INTEGER_RE = re.compile(r"^-?\d+$")
# 1 234 567,89 · 1234.5 · -12,5 · 0,5 (разделители тысяч: обычный, неразрывный, узкий пробел)
DECIMAL_RE = re.compile(r"^-?([1-9]\d{0,2}([ \u00a0\u202f]\d{3})+|0|[1-9]\d*)([.,]\d+)?$")
DATE_RE = re.compile(r"^(\d{2})\.(\d{2})\.(\d{4})$")


class ConversionError(Exception):
    """Ошибка, понятная пользователю: текст показывается как есть."""


@dataclass
class Table:
    pages: list[int]
    bbox: pymupdf.Rect
    rows: list[list[str]] = field(default_factory=list)

    @property
    def col_count(self):
        return len(self.rows[0]) if self.rows else 0

    @property
    def title(self):
        first, last = self.pages[0], self.pages[-1]
        return f"стр. {first}" if first == last else f"стр. {first}–{last}"


@dataclass
class Result:
    xlsx: bytes
    pages: int
    tables: int
    rows: int
    mode: str  # lines | text | plain


def clean(text):
    """Убирает мягкие переносы, переносы со знаком U+2010 и склеивает строки ячейки."""
    text = text.replace("\u00ad", "").replace("\u00a0", " ")
    text = ILLEGAL_CHARACTERS_RE.sub("", text)  # управляющие символы Excel не принимает
    text = re.sub(r"\u2010\s*\n\s*", "", text)
    text = re.sub(r"\s*\n\s*", " ", text)
    return re.sub(r"[ \t]{2,}", " ", text).strip()


def page_spans(page):
    """Текстовые фрагменты страницы с координатами — один проход вместо поиска по каждой ячейке."""
    spans = []
    for block in page.get_text("rawdict")["blocks"]:
        for line in block.get("lines", []):
            line_id = (block["number"], round(line["bbox"][1]))
            chars = [(i, ch) for i, span in enumerate(line["spans"]) for ch in span["chars"]]
            texts = [[] for _ in line["spans"]]
            for n, (i, ch) in enumerate(chars):
                c = ch["c"]
                if c == "\u00ad":
                    # Chrome ставит невидимые мягкие переносы (следующая буква не сдвигается),
                    # а часть генераторов кодирует этим символом обычный видимый дефис
                    nxt = chars[n + 1][1] if n + 1 < len(chars) else None
                    if nxt is None or nxt["bbox"][0] - ch["bbox"][0] < 0.5:
                        continue
                    c = "-"
                texts[i].append(c)
            for span, text in zip(line["spans"], texts):
                spans.append((pymupdf.Rect(span["bbox"]), "".join(text), line_id))
    return spans


def cell_text(spans, rect, exclude):
    if rect is None:
        return ""
    rect = pymupdf.Rect(rect)
    # если внутри ячейки лежит вложенная таблица — берём только текст над ней
    for inner in exclude:
        if rect.intersects(inner) and inner.y0 > rect.y0:
            rect.y1 = min(rect.y1, inner.y0)
    parts, prev_line = [], None
    for bbox, text, line_id in spans:
        center = pymupdf.Point((bbox.x0 + bbox.x1) / 2, (bbox.y0 + bbox.y1) / 2)
        if center not in rect:
            continue
        if prev_line is not None and line_id != prev_line:
            parts.append("\n")
        parts.append(text)
        prev_line = line_id
    return clean("".join(parts))


def find_tables(doc, strategy):
    tables = []
    for page in doc:
        found = page.find_tables(strategy=strategy).tables
        if not found:
            continue
        spans = page_spans(page)
        rects = [pymupdf.Rect(t.bbox) for t in found]
        for table, rect in zip(found, rects):
            inner = [r for r in rects if r is not rect and r in rect]
            rows = [[cell_text(spans, c, inner) for c in row.cells] for row in table.rows]
            rows = [r for r in rows if any(r)]
            if rows:
                tables.append(Table(pages=[page.number + 1], bbox=rect, rows=rows))
    return tables


def merge_continued(tables):
    """Склеивает таблицу, разорванную между страницами: та же ширина и колонки, идёт сразу следом."""
    merged = []
    for table in tables:
        prev = merged[-1] if merged else None
        if (
            prev
            and table.pages[0] == prev.pages[-1] + 1
            and table.col_count == prev.col_count
            and abs(table.bbox.x0 - prev.bbox.x0) < 5
            and abs(table.bbox.x1 - prev.bbox.x1) < 5
        ):
            rows = table.rows[1:] if table.rows[0] == prev.rows[0] else table.rows  # повтор шапки
            prev.rows.extend(rows)
            prev.pages.append(table.pages[0])
            prev.bbox = table.bbox
        else:
            merged.append(table)
    return merged


def to_value(text):
    if not text:
        return None
    if INTEGER_RE.match(text):
        digits = text.lstrip("-")
        # артикулы с ведущим нулём и длинные номера (ИНН, счета) — текстом, иначе Excel их испортит
        if len(digits) >= 10 or (len(digits) > 1 and digits[0] == "0"):
            return text
        return int(text)
    if DECIMAL_RE.match(text):
        return float(re.sub(r"[ \u00a0\u202f]", "", text).replace(",", "."))
    if match := DATE_RE.match(text):
        day, month, year = map(int, match.groups())
        try:
            return datetime(year, month, day)
        except ValueError:
            return text
    return text


def write_rows(ws, rows, start_row):
    bold = Font(bold=True)
    for offset, row in enumerate(rows):
        for col, text in enumerate(row, start=1):
            cell = ws.cell(row=start_row + offset, column=col, value=text if offset == 0 else to_value(text))
            if offset == 0:
                cell.font = bold
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            elif isinstance(cell.value, datetime):
                cell.number_format = "DD.MM.YYYY"
    return start_row + len(rows)


def autosize(ws):
    widths = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None:
                length = 10 if isinstance(cell.value, datetime) else len(str(cell.value))
                widths[cell.column] = max(widths.get(cell.column, 0), length)
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = min(max(width + 2, 8), 60)


def build_workbook(tables, single_sheet):
    wb = Workbook()
    if single_sheet:
        ws = wb.active
        ws.title = "Таблицы"
        row = 1
        for n, table in enumerate(tables, start=1):
            ws.cell(row=row, column=1, value=f"Таблица {n} · {table.title}").font = Font(bold=True, size=12)
            row = write_rows(ws, table.rows, row + 1) + 1
        autosize(ws)
    else:
        wb.remove(wb.active)
        for n, table in enumerate(tables, start=1):
            ws = wb.create_sheet(f"Таблица {n} ({table.title})"[:31])
            write_rows(ws, table.rows, 1)
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            autosize(ws)
    return wb


def build_text_workbook(doc):
    wb = Workbook()
    ws = wb.active
    ws.title = "Текст"
    ws.append(["Стр.", "Текст"])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for page in doc:
        for line in page.get_text().split("\n"):
            if line := clean(line):
                ws.append([page.number + 1, line])
    ws.column_dimensions["B"].width = 120
    return wb


def convert(data, *, single_sheet=False, max_pages=None):
    try:
        doc = pymupdf.open(stream=data, filetype="pdf")
    except Exception as exc:
        raise ConversionError("Файл повреждён или это не PDF") from exc

    with doc:
        if doc.needs_pass:
            raise ConversionError("PDF защищён паролем")
        if doc.page_count == 0:
            raise ConversionError("В PDF нет страниц")
        if max_pages and doc.page_count > max_pages:
            raise ConversionError(f"Слишком много страниц: {doc.page_count}, максимум {max_pages}")
        if not any(page.get_text().strip() for page in doc):
            raise ConversionError("В PDF нет текстового слоя — похоже на скан. Распознавание (OCR) не поддерживается")

        # таблицы с линиями; если их нет совсем — пробуем выравнивание текста (таблицы без рамок)
        mode = "lines"
        tables = merge_continued(find_tables(doc, "lines"))
        if not tables:
            mode = "text"
            tables = merge_continued(find_tables(doc, "text"))
        if tables:
            wb = build_workbook(tables, single_sheet)
        else:
            mode = "plain"
            wb = build_text_workbook(doc)

        buffer = io.BytesIO()
        wb.save(buffer)
        return Result(
            xlsx=buffer.getvalue(),
            pages=doc.page_count,
            tables=len(tables),
            rows=sum(len(t.rows) for t in tables),
            mode=mode,
        )


def main():
    parser = argparse.ArgumentParser(description="PDF → Excel")
    parser.add_argument("files", nargs="+", type=Path)
    parser.add_argument("-o", "--output", type=Path, help="папка для .xlsx (по умолчанию рядом с PDF)")
    parser.add_argument("--single-sheet", action="store_true", help="все таблицы на одном листе")
    args = parser.parse_args()

    failed = 0
    for pdf in args.files:
        target = (args.output or pdf.parent) / f"{pdf.stem}.xlsx"
        try:
            result = convert(pdf.read_bytes(), single_sheet=args.single_sheet)
        except ConversionError as exc:
            print(f"{pdf.name}: {exc}", file=sys.stderr)
            failed += 1
            continue
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_bytes(result.xlsx)
        print(f"{pdf.name} -> {target}: таблиц {result.tables}, строк {result.rows} ({result.mode})")
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
